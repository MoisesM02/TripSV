from pathlib import Path
from datetime import date

from PIL import Image, ImageDraw, ImageEnhance, ImageFont, ImageOps
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\moise\Documents\UDB\2026\Ciclo 2\DES\Proyecto Cátedra\TripSV")
OUT = ROOT / "Documentacion"
ASSETS = OUT / "assets"
APP = ROOT / "TripSV"

BLUE = "0D6EFD"
RED = "E5394F"
DARK = "1F2937"
MUTED = "667085"
PALE = "EEF4FF"
GRID = "D9D9D9"


def pil_color(value):
    if value.startswith("#") or value.lower() in {"white", "black", "gray", "grey"}:
        return value
    return f"#{value}"


def font(size, bold=False, color=DARK, name="Aptos"):
    f = ImageFont.truetype(r"C:\Windows\Fonts\segoeui.ttf", size)
    if bold:
        f = ImageFont.truetype(r"C:\Windows\Fonts\segoeuib.ttf", size)
    return f


def fit_cover(path, size):
    try:
        img = Image.open(path).convert("RGB")
    except Exception:
        img = Image.new("RGB", size, "#DDE5F2")
    return ImageOps.fit(img, size, method=Image.Resampling.LANCZOS)


def rounded(draw, box, fill, outline=None, radius=16, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=pil_color(fill) if isinstance(fill, str) else fill,
                           outline=pil_color(outline) if isinstance(outline, str) else outline, width=width)


def make_figures():
    ASSETS.mkdir(parents=True, exist_ok=True)
    tazumal = APP / "wwwroot" / "img" / "tazumal.jpg"
    tunco = APP / "wwwroot" / "img" / "eltincoo.jpg"
    zonte = APP / "wwwroot" / "img" / "zonte.jpg"
    guirola = ASSETS / "mansion-guirola.png"

    # portada
    im = Image.new("RGB", (1200, 680), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1200, 56), fill="#1677F2")
    d.text((24, 15), "El Salvador", fill="white", font=font(24, True))
    d.text((190, 18), "Categorías", fill="#DCE8FF", font=font(17))
    d.text((300, 18), "Sitios", fill="#DCE8FF", font=font(17))
    rounded(d, (800, 12, 1040, 44), "white", radius=7)
    d.text((820, 19), "Buscar destinos", fill="#98A2B3", font=font(15))
    rounded(d, (1050, 12, 1140, 44), "#1677F2", outline="white", radius=7)
    d.text((1070, 19), "Buscar", fill="white", font=font(15))
    hero = fit_cover(tazumal, (1200, 420))
    hero = ImageEnhance.Brightness(hero).enhance(0.78)
    im.paste(hero, (0, 56))
    d.text((530, 430), "Tazumal", fill="white", font=font(29, True))
    rounded(d, (72, 445, 1128, 570), "white", outline="#D0D5DD", radius=14)
    d.text((98, 470), "¿A dónde quiere ir?", fill="#101828", font=font(23, True))
    rounded(d, (98, 505, 880, 550), "white", outline="#D0D5DD", radius=8)
    d.text((118, 517), "Playas, lagos, volcanes, pueblos...", fill="#98A2B3", font=font(16))
    rounded(d, (894, 505, 1102, 550), RED, radius=8)
    d.text((970, 518), "Buscar", fill="white", font=font(16, True))
    d.text((72, 610), "Categorías turísticas", fill="#101828", font=font(22, True))
    im.save(ASSETS / "fig_portada.png")

    # catalogo
    im = Image.new("RGB", (1200, 720), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1200, 56), fill="#1677F2")
    d.text((24, 15), "El Salvador", fill="white", font=font(24, True))
    d.text((190, 18), "Categorías", fill="#DCE8FF", font=font(17))
    d.text((300, 18), "Sitios", fill="#DCE8FF", font=font(17))
    d.text((72, 88), "Catálogo de sitios turísticos", fill="#101828", font=font(30, True))
    rounded(d, (72, 138, 1128, 268), "#F8FAFC", outline="#D0D5DD", radius=12)
    labels = [(90, "Buscar"), (405, "Categoría"), (620, "Departamento"), (835, "Calificación"), (1020, "Ordenar por")]
    for x, label in labels:
        d.text((x, 156), label, fill="#475467", font=font(15))
    rounded(d, (90, 183, 375, 225), "white", outline="#D0D5DD", radius=7)
    d.text((107, 195), "Nombre, lugar o tipo de destino", fill="#98A2B3", font=font(13))
    for x, val in [(405, "Todas"), (620, "Todos"), (835, "Cualquiera"), (1020, "Mejor calificados")]:
        rounded(d, (x, 183, x + 180, 225), "white", outline="#D0D5DD", radius=7)
        d.text((x + 14, 195), val, fill="#475467", font=font(14))
    rounded(d, (90, 235, 180, 265), RED, radius=7)
    d.text((116, 242), "Buscar", fill="white", font=font(14, True))
    d.text((72, 302), "27 destino(s) en el catálogo", fill="#667085", font=font(15))
    cards = [(tazumal, "El Tazumal", "Santa Ana"), (zonte, "El Zonte", "La libertad"), (guirola, "Mansión Guirola", "La Libertad")]
    for idx, (pic, title, place) in enumerate(cards):
        x = 74 + idx * 365
        rounded(d, (x, 344, x + 330, 690), "white", outline="#D0D5DD", radius=10)
        im.paste(fit_cover(pic, (330, 170)), (x, 344))
        d.text((x + 16, 530), title, fill="#101828", font=font(18, True))
        d.text((x + 16, 558), place, fill="#667085", font=font(15))
        rounded(d, (x + 16, 586, x + 170, 616), BLUE, radius=7)
        d.text((x + 28, 593), "Sitio turístico", fill="white", font=font(13, True))
        d.text((x + 16, 640), "★★★★★  4.5/5", fill="#F59E0B", font=font(16, True))
    im.save(ASSETS / "fig_catalogo.png")

    # detalle
    im = Image.new("RGB", (1200, 720), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1200, 56), fill="#1677F2")
    d.text((24, 15), "El Salvador", fill="white", font=font(24, True))
    d.text((190, 18), "Categorías", fill="#DCE8FF", font=font(17))
    d.text((300, 18), "Sitios", fill="#DCE8FF", font=font(17))
    d.text((72, 86), "El Tazumal, Santa Ana", fill="#101828", font=font(30, True))
    d.text((72, 130), "El sitio arqueológico Tazumal está ubicado en Chalchuapa.", fill="#667085", font=font(16))
    im.paste(fit_cover(tazumal, (720, 440)), (72, 175))
    rounded(d, (825, 175, 1128, 335), "white", outline="#D0D5DD", radius=10)
    d.text((850, 200), "Favoritos", fill="#101828", font=font(21, True))
    d.text((850, 238), "1 persona guardó este destino.", fill="#667085", font=font(14))
    rounded(d, (850, 270, 1102, 312), RED, radius=7)
    d.text((900, 282), "♥ Guardar", fill="white", font=font(15, True))
    rounded(d, (825, 360, 1128, 540), "white", outline="#D0D5DD", radius=10)
    d.text((850, 385), "Calificación", fill="#101828", font=font(21, True))
    d.text((850, 425), "★★★★★  4.5/5", fill="#F59E0B", font=font(19, True))
    d.text((850, 470), "Publicar reseña y responder", fill="#475467", font=font(14))
    d.text((72, 650), "Reseñas de la comunidad", fill="#101828", font=font(22, True))
    im.save(ASSETS / "fig_detalle.png")

    # itinerario
    im = Image.new("RGB", (1200, 720), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1200, 56), fill="#1677F2")
    d.text((24, 15), "El Salvador", fill="white", font=font(24, True))
    d.text((190, 18), "Mis itinerarios", fill="#DCE8FF", font=font(17))
    d.text((72, 92), "Ruta de revisión", fill="#101828", font=font(30, True))
    d.text((72, 135), "01/10/2026 · 1 día(s) · 2 destino(s)", fill="#667085", font=font(16))
    rounded(d, (72, 178, 1128, 280), "#F8FAFC", outline="#D0D5DD", radius=10)
    d.text((95, 202), "Prueba funcional del planificador.", fill="#475467", font=font(16))
    rounded(d, (72, 310, 1128, 650), "white", outline="#D0D5DD", radius=10)
    d.rectangle((72, 310, 1128, 365), fill="#EAF0FA")
    d.text((94, 326), "Día 1", fill="#101828", font=font(19, True))
    d.text((180, 328), "Jueves 1 de octubre", fill="#475467", font=font(16))
    rows = [(tazumal, "El Tazumal", "Santa Ana · Sitios arqueológicos"), (zonte, "El Zonte", "La libertad · Playa")]
    for i, (pic, title, meta) in enumerate(rows):
        y = 382 + i * 115
        im.paste(fit_cover(pic, (90, 80)), (96, y))
        d.ellipse((205, y + 10, 239, y + 44), fill=pil_color(BLUE))
        d.text((216, y + 14), str(i + 1), fill="white", font=font(15, True))
        d.text((265, y + 5), title, fill="#1769D4", font=font(18, True))
        d.text((265, y + 38), meta, fill="#667085", font=font(14))
        d.text((1050, y + 25), "Quitar", fill=pil_color(RED), font=font(14, True))
    rounded(d, (96, 608, 650, 642), "white", outline="#D0D5DD", radius=7)
    d.text((112, 615), "Agregar un destino a este día...", fill="#98A2B3", font=font(14))
    rounded(d, (990, 608, 1100, 642), RED, radius=7)
    d.text((1015, 615), "Agregar", fill="white", font=font(14, True))
    im.save(ASSETS / "fig_itinerario.png")

    # administración
    im = Image.new("RGB", (1200, 720), "white")
    d = ImageDraw.Draw(im)
    d.rectangle((0, 0, 1200, 56), fill="#1677F2")
    d.text((24, 15), "El Salvador", fill="white", font=font(24, True))
    d.text((190, 18), "Administrar", fill="#DCE8FF", font=font(17))
    d.text((72, 92), "Moderación de comentarios", fill="#101828", font=font(29, True))
    for x, val, lab in [(72, "83", "Visibles"), (220, "0", "Ocultos")]:
        rounded(d, (x, 145, x + 120, 220), "#F8FAFC", outline="#D0D5DD", radius=8)
        d.text((x + 42, 155), val, fill=pil_color(BLUE), font=font(24, True))
        d.text((x + 20, 192), lab, fill="#667085", font=font(13))
    rounded(d, (380, 145, 780, 185), "white", outline="#D0D5DD", radius=7)
    d.text((400, 156), "Buscar por comentario, usuario o sitio", fill="#98A2B3", font=font(13))
    rounded(d, (800, 145, 900, 185), BLUE, radius=7)
    d.text((827, 156), "Filtrar", fill="white", font=font(14, True))
    rounded(d, (72, 250, 1128, 665), "white", outline="#D0D5DD", radius=8)
    d.rectangle((72, 250, 1128, 300), fill="#20242C")
    headers = [(92, "ID"), (155, "Usuario"), (285, "Sitio"), (470, "Comentario"), (835, "Estado"), (1000, "Acciones")]
    for x, label in headers:
        d.text((x, 267), label, fill="white", font=font(14, True))
    rows = [("163", "Lucas", "Costa del Sol", "Me ha servido mucho...", "Visible"), ("161", "Eduardo", "Guatajiagua", "Excellent place.", "Visible"), ("143", "administrador", "El Tazumal", "Me encanta la historia...", "Visible")]
    for i, row in enumerate(rows):
        y = 300 + i * 110
        d.rectangle((72, y, 1128, y + 108), fill="#F8FAFC" if i % 2 == 0 else "white")
        for x, text in [(92, row[0]), (155, row[1]), (285, row[2]), (470, row[3]), (835, row[4])]:
            d.text((x, y + 25), text, fill="#475467", font=font(14))
        d.text((1000, y + 25), "Ocultar", fill=pil_color(RED), font=font(14, True))
        d.line((72, y + 108, 1128, y + 108), fill="#D9E0EA", width=1)
    im.save(ASSETS / "fig_admin.png")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=GRID, size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_widths(table, widths):
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
        for row in table.rows:
            row.cells[idx].width = Inches(width)


def set_run_font(run, name="Aptos", size=None, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for style_name, size in (("Title", 28), ("Heading 1", 19), ("Heading 2", 14), ("Heading 3", 11.5)):
        st = styles[style_name]
        st.font.name = "Aptos Display" if style_name in ("Title", "Heading 1") else "Aptos"
        st._element.rPr.rFonts.set(qn("w:ascii"), st.font.name)
        st._element.rPr.rFonts.set(qn("w:hAnsi"), st.font.name)
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string("000000")
        st.paragraph_format.space_before = Pt(14 if style_name != "Heading 3" else 8)
        st.paragraph_format.space_after = Pt(5)
        st.paragraph_format.keep_with_next = True


def add_footer(doc, label):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"Trips SV | {label} | 29 de septiembre de 2026")
        set_run_font(r, size=8.5, color=MUTED)


def add_cover(doc, title, subtitle, audience, scope):
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run(title)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    set_run_font(r, size=15, color=BLUE, bold=True)
    p.paragraph_format.space_after = Pt(18)
    table = doc.add_table(rows=4, cols=2)
    table.autofit = False
    cover_widths = [1.5, 5.2]
    rows = [("Sistema", "Trips SV - CMS de turismo de El Salvador"), ("Audiencia", audience), ("Alcance", scope), ("Versión", "1.0 | Revisión funcional y técnica local")]
    for i, (a, b) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        set_cell_margins(c0); set_cell_margins(c1)
        set_cell_shading(c0, BLUE); set_cell_shading(c1, PALE if i % 2 == 0 else "FFFFFF")
        c0.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c1.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r0 = c0.paragraphs[0].add_run(a); set_run_font(r0, color="FFFFFF", bold=True, size=10)
        r1 = c1.paragraphs[0].add_run(b); set_run_font(r1, size=10)
    set_table_widths(table, cover_widths)
    set_table_borders(table)
    doc.add_paragraph()
    p = doc.add_paragraph()
    r = p.add_run("Documento preparado a partir de la revisión del código fuente y del recorrido de la aplicación en modo Development. Las figuras muestran referencias visuales de las pantallas observadas y de sus controles principales.")
    set_run_font(r, size=10.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_toc(doc, entries):
    doc.add_heading("Contenido", level=1)
    for num, title in entries:
        p = doc.add_paragraph(style="Normal")
        p.paragraph_format.left_indent = Inches(0.2)
        r = p.add_run(f"{num}. {title}")
        set_run_font(r, bold=True, size=10.5)
    doc.add_page_break()


def add_para(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead); set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_lead):]); set_run_font(r)
    else:
        r = p.add_run(text); set_run_font(r)
    return p


def add_bullets(doc, items, level=0):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.25 + level * 0.2)
        r = p.add_run(item); set_run_font(r, size=10.3)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        r = p.add_run(f"{index}. "); set_run_font(r, size=10.3, bold=True)
        r = p.add_run(item); set_run_font(r, size=10.3)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, BLUE); set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        r = cell.paragraphs[0].add_run(h); set_run_font(r, color="FFFFFF", bold=True, size=9.5)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for cidx, value in enumerate(row):
            cell = cells[cidx]
            set_cell_shading(cell, PALE if ridx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            r = cell.paragraphs[0].add_run(str(value)); set_run_font(r, size=9.3)
    if widths:
        set_table_widths(table, widths)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, filename, caption, width=6.7):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cp.add_run(caption); set_run_font(r, size=8.8, color=MUTED, italic=True)


def add_capture(doc, filename, caption, width=6.35):
    path = OUT / "fuentes" / "capturas" / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Inches(width))
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.keep_with_next = False
    r = cp.add_run(caption); set_run_font(r, size=8.8, color=MUTED, italic=True)


def add_code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F3F4F6"); pPr.append(shd)
    r = p.add_run(text); set_run_font(r, name="Consolas", size=8.8, color="111827")


def build_programmer():
    doc = Document(); configure_doc(doc)
    add_cover(doc, "Manual del Programador", "Arquitectura, configuración y despliegue de Trips SV", "Desarrolladores, administradores técnicos y responsables de despliegue.", "Describe la arquitectura MVC por capas, el acceso seguro a SQL Server, la interfaz, las pruebas y los procedimientos para ejecutar, publicar y mantener el sistema.")
    add_toc(doc, [
        ("1", "Resumen técnico"), ("2", "Arquitectura del sistema"), ("3", "Tecnologías utilizadas"),
        ("4", "Estructura de la base de código"), ("5", "Modelo de datos"), ("6", "Acceso a datos y persistencia"),
        ("7", "Interfaces de usuario, responsividad y accesibilidad"), ("8", "Seguridad y manejo de errores"),
        ("9", "Configuración del entorno"), ("10", "Despliegue"), ("11", "Pruebas automatizadas"),
        ("12", "Mantenimiento y extensión")
    ])

    doc.add_heading("1 Resumen técnico", level=1)
    add_para(doc, "Trips SV es un CMS web para la publicación y exploración de destinos turísticos de El Salvador. La versión revisada está construida con ASP.NET Core MVC sobre .NET 10, Entity Framework Core Code First, SQL Server y ASP.NET Core Identity. El catálogo cargado por el script de datos contiene 9 categorías, 27 sitios y comentarios de la comunidad.")
    add_table(doc, ["Elemento", "Implementación verificada"], [
        ("Framework", ".NET 10 / ASP.NET Core MVC con Razor"),
        ("Persistencia", "Entity Framework Core 10 con proveedor SQL Server y migraciones Code First"),
        ("Autenticación", "ASP.NET Core Identity con roles Administrador y Usuario"),
        ("Interfaz", "Bootstrap local, Bootstrap Icons, CSS propio, jQuery Validation Unobtrusive y CKEditor"),
        ("Seguridad de contenido", "HtmlSanitizer para el contenido enriquecido y validación de imágenes"),
        ("Pruebas", "Proyecto TripSV.Pruebas con xUnit y EF Core SQLite en memoria; el repositorio declara 80 casos")
    ], [1.65, 5.05])

    doc.add_heading("2 Arquitectura del sistema", level=1)
    add_para(doc, "Tipo de arquitectura: MVC con N-capas lógicas dentro de un único proyecto web. MVC separa Modelo, Vista y Controlador; la separación por capas agrega una capa de servicios de negocio entre los controladores y el acceso a datos. No es MVVM: las vistas Razor se renderizan en el servidor y los ViewModels transportan datos y validación, pero no existe un enlace bidireccional reactivo propio de una aplicación SPA.")
    add_table(doc, ["Capa", "Responsabilidad", "Componentes principales"], [
        ("Presentación", "Renderiza páginas, formularios y mensajes; aplica validación de entrada y autorización.", "Controllers, Areas/Administracion, Views Razor, ViewModels, wwwroot"),
        ("Negocio", "Centraliza reglas de búsqueda, comentarios, calificaciones, favoritos, itinerarios y resultados de operación.", "Interfaces e implementaciones en Servicios/"),
        ("Datos", "Mapea entidades, relaciones, índices y restricciones; persiste cambios y ejecuta migraciones.", "Modelos/, ContextoTripSV, Configuraciones/, Migraciones/"),
        ("Infraestructura", "Compone el contenedor, Identity, middleware, archivos estáticos y configuración.", "Program.cs, appsettings.json, ASP.NET Core Identity")
    ], [1.15, 2.8, 2.75])
    add_code(doc, "Solicitud HTTP -> Controller -> Servicio (interfaz) -> ContextoTripSV -> SQL Server\n                         -> ViewModel -> Razor View -> Respuesta HTML")
    doc.add_heading("2.1 Justificación de la elección", level=2)
    add_bullets(doc, [
        "Renderizado del lado del servidor: el catálogo y las páginas públicas reciben HTML completo, con una carga de JavaScript reducida y una navegación apropiada para un sitio de contenidos.",
        "Formularios con validación integrada: MVC combina model binding, DataAnnotations, validación unobtrusive y tokens antifalsificación para las operaciones de cuenta, reseñas y administración.",
        "Reglas testeables: los controladores son delgados y los servicios pueden probarse con un ContextoTripSV sobre SQLite en memoria sin levantar el servidor web.",
        "Separación de responsabilidades: la persistencia queda concentrada en EF Core, las reglas en Servicios y el módulo administrativo en un Area protegida por rol.",
        "Proporción con el proyecto: un monolito modular evita el costo de varios despliegues y comunicaciones de microservicios para un dominio pequeño/mediano, pero mantiene límites claros para evolucionar el sistema."
    ])
    doc.add_heading("2.2 Flujo de una operación", level=2)
    add_numbered(doc, [
        "El navegador envía una solicitud, por ejemplo POST /Puntuaciones/Calificar, con el token antifalsificación.",
        "El controlador comprueba la sesión mediante [Authorize] y valida el modelo recibido.",
        "El controlador invoca IPuntuacionesServicio; el servicio comprueba el rango 1–5, crea o actualiza la calificación y recalcula el promedio.",
        "El servicio devuelve Resultado.Ok o Resultado.Error; el controlador guarda el mensaje en TempData y redirige al detalle usando Post/Redirect/Get."
    ])
    add_table(doc, ["Patrón o principio", "Aplicación en Trips SV"], [
        ("Inyección de dependencias", "Program.cs registra interfaces; controladores y servicios reciben sus dependencias por constructor."),
        ("Unidad de trabajo", "ContextoTripSV y sus DbSet coordinan las lecturas y escrituras de una solicitud."),
        ("ViewModel", "Las pantallas que combinan datos usan modelos de vista específicos, sin depender siempre de las entidades."),
        ("Result", "Resultado comunica errores esperables sin usar excepciones como flujo normal."),
        ("SOLID", "Interfaces pequeñas por módulo; validación de imágenes y saneamiento HTML en servicios separados.")
    ], [1.65, 5.05])

    doc.add_heading("3 Tecnologías utilizadas", level=1)
    add_table(doc, ["Componente", "Tecnología", "Uso"], [
        ("Plataforma", ".NET 10 / ASP.NET Core", "Servidor web y ciclo de vida de la aplicación."),
        ("Presentación", "MVC, Razor, Bootstrap 5, Bootstrap Icons", "Vistas, navegación, grid responsive e iconografía."),
        ("Datos", "EF Core 10 / SQL Server", "Mapeo ORM, consultas LINQ, restricciones y migraciones."),
        ("Identidad", "ASP.NET Core Identity", "Usuarios, contraseñas, cookies, roles y tokens."),
        ("Contenido", "HtmlSanitizer 9.x / CKEditor", "Edición y saneamiento de información enriquecida."),
        ("Cliente", "jQuery Validation Unobtrusive", "Validación de formularios en el navegador."),
        ("Pruebas", "xUnit / EF Core SQLite", "Pruebas de servicios, validación, saneamiento y autorización.")
    ], [1.3, 2.15, 3.25])

    doc.add_heading("4 Estructura de la base de código", level=1)
    add_code(doc, "TripSV.slnx\n├── TripSV/\n│   ├── Program.cs                 Composición, middleware y rutas\n│   ├── Modelos/                   Entidades y roles\n│   ├── Datos/                     DbContext, Fluent API, migraciones y sembrador\n│   ├── Servicios/                 Interfaces y reglas de negocio\n│   ├── ViewModels/                Modelos para formularios y páginas\n│   ├── Controllers/               Flujos públicos y de usuario\n│   ├── Areas/Administracion/      CRUD y moderación protegidos\n│   ├── Views/                     Vistas Razor y parciales\n│   └── wwwroot/                   CSS, JS, imágenes y librerías\n├── TripSV.Pruebas/                Pruebas automatizadas\n├── Herramientas/                  Utilidades de datos/documentación\n├── Documentacion/                 Manuales, fuentes y capturas\n└── Ejecutable/                    Publicación Windows x64")
    add_table(doc, ["Módulo", "Responsabilidad"], [
        ("Controllers", "Inicio, categorías, catálogo, cuenta, comentarios, puntuaciones, favoritos e itinerarios."),
        ("Areas/Administracion", "Categorías, sitios y moderación; las acciones exigen el rol Administrador."),
        ("Servicios", "Reglas del dominio, validaciones, consultas y guardado con manejo de conflictos."),
        ("Datos", "ContextoTripSV, configuraciones IEntityTypeConfiguration, migraciones y seed de Identity."),
        ("Views y ViewModels", "Interfaz pública, formularios, mensajes, tablas, tarjetas y pantallas administrativas."),
        ("TripSV.Pruebas", "Casos de xUnit para servicios, validación de imágenes, saneamiento de HTML y seguridad.")
    ], [2.0, 4.7])
    add_para(doc, "El menú Administrar se oculta para otros roles, pero la garantía real está en la autorización del servidor. Las vistas parciales _TarjetaSitio, _TarjetaCategoria, _Mensajes y los componentes de estrellas permiten una interfaz consistente.")

    doc.add_heading("5 Modelo de datos", level=1)
    add_para(doc, "ContextoTripSV hereda de IdentityDbContext<Usuario>, aplica las configuraciones de todo el ensamblado con ApplyConfigurationsFromAssembly y expone DbSet para el dominio de turismo.")
    add_table(doc, ["Tabla / entidad", "Reglas y relaciones importantes"], [
        ("categorias / Categoria", "Nombre único; una categoría tiene muchos sitios y no se elimina si aún tiene destinos."),
        ("sitios / Sitio", "Nombre único; pertenece a una categoría; guarda imagen, información enriquecida, promedio y total de puntuaciones."),
        ("comentarios / Comentario", "Reseñas y respuestas; puede ocultarse; respuestas enlazadas mediante respuesta_a_id."),
        ("puntuaciones / Puntuacion", "Valor de 1 a 5; índice único por sitio y usuario; CHECK de base de datos."),
        ("favoritos / Favorito", "Índice único por usuario y sitio; borrado en cascada con usuario o sitio."),
        ("itinerarios / Itinerario", "Pertenece a un usuario; CHECK de 1 a 15 días."),
        ("visitas / Visita", "Destino por día y orden; índice único por itinerario, día y sitio."),
        ("Identity", "usuarios, roles, usuarios_roles y tablas de claims, logins y tokens con nombres en español.")
    ], [1.8, 4.9])
    add_table(doc, ["Al eliminar", "Comportamiento"], [
        ("Destino", "Cascada para comentarios, puntuaciones, favoritos y visitas."),
        ("Usuario", "Cascada para favoritos e itinerarios; comentarios y calificaciones conservan el nombre del autor."),
        ("Itinerario", "Cascada para sus visitas."),
        ("Comentario", "El servicio elimina primero sus respuestas."),
        ("Categoría", "Restrict: se impide mientras existan sitios asociados.")
    ], [1.8, 4.9])

    doc.add_heading("6 Acceso a datos y persistencia", level=1)
    doc.add_heading("6.1 Entity Framework Core Code First", level=2)
    add_para(doc, "El esquema se define en C# mediante entidades y clases de configuración Fluent API. ContextoTripSV traduce las consultas LINQ a SQL Server, mientras que las migraciones materializan cambios de esquema. SembradorDatos ejecuta Database.MigrateAsync al iniciar y crea los roles y usuarios iniciales si no existen.")
    add_code(doc, "dotnet ef migrations add NombreDeLaMigracion --project TripSV --output-dir Datos/Migraciones\ndotnet ef database update --project TripSV")
    add_para(doc, "Las clases de Configuraciones definen nombres snake_case, tipos de columna, longitudes, índices únicos, CHECK constraints y comportamientos de borrado. ExtensionesContexto.GuardarSinConflictoAsync encapsula SaveChangesAsync, limpia el ChangeTracker cuando ocurre un DbUpdateException y devuelve un resultado de conflicto para que el servicio informe al usuario.")
    doc.add_heading("6.2 Cadena de conexión segura", level=2)
    add_para(doc, "Program.cs obtiene la clave ConnectionStrings:ConexionTripSV mediante GetConnectionString y falla temprano si no existe. La configuración local actual usa autenticación integrada de Windows contra LocalDB; no incluye usuario ni contraseña.")
    add_code(doc, '"ConexionTripSV": "Server=(localdb)\\mssqllocaldb;Database=TripSV;Trusted_Connection=True;TrustServerCertificate=True;MultipleActiveResultSets=True"')
    add_table(doc, ["Entorno", "Mecanismo recomendado", "Ejemplo"], [
        ("Desarrollo", "User Secrets fuera del repositorio", "dotnet user-secrets set \"ConnectionStrings:ConexionTripSV\" \"Server=...;\" --project TripSV"),
        ("CI/CD", "Variable protegida del pipeline", "ConnectionStrings__ConexionTripSV"),
        ("Producción", "Secret manager o variable de entorno", "Server=...;User Id=...;Password=...;Encrypt=True")
    ], [1.05, 2.05, 3.6])
    add_bullets(doc, [
        "Nunca versionar contraseñas de SQL Server, tokens o credenciales de demostración.",
        "TrustServerCertificate=True es aceptable solo para desarrollo local; en producción usar Encrypt=True y un certificado válido.",
        "El usuario de producción debe tener permisos mínimos y la cadena debe inyectarse desde el ambiente, no editarse en el binario publicado."
    ])
    doc.add_heading("6.3 Uso eficiente de la persistencia", level=2)
    add_bullets(doc, [
        "Los listados usan proyecciones que omiten la columna binaria de imagen; SinImagen selecciona solo los campos necesarios y la imagen se obtiene por /Sitios/Imagen/{id}.",
        "Las imágenes del endpoint dedicado tienen ResponseCache de una hora; los archivos estáticos tienen Cache-Control de siete días.",
        "Los conteos, filtros y ordenamientos se ejecutan en SQL mediante LINQ; la búsqueda admite hasta cinco términos y usa una collation tolerante a tildes y mayúsculas.",
        "Los índices únicos también hacen cumplir reglas de negocio para nombres, favoritos, puntuaciones y visitas.",
        "Las lecturas que no necesitan seguimiento usan proyecciones o AsNoTracking; las operaciones de actualización mantienen el seguimiento solo de la entidad que se modifica.",
        "La calificación promedio y el total de votos se almacenan en Sitio y se recalculan únicamente cuando cambia una puntuación."
    ])
    doc.add_heading("6.4 Datos iniciales", level=2)
    add_para(doc, "El catálogo no se carga desde SembradorDatos. TripSV/Datos/Migracion/datos-mysql.sql contiene categorías, destinos, imágenes, reseñas y puntuaciones, y debe ejecutarse una vez después de crear el esquema:")
    add_code(doc, 'sqlcmd -S "(localdb)\\MSSQLLocalDB" -d TripSV -E -i "TripSV\\Datos\\Migracion\\datos-mysql.sql" -b')

    doc.add_heading("7 Interfaces de usuario, responsividad y accesibilidad", level=1)
    add_para(doc, "La interfaz es una aplicación server-rendered con Razor, Bootstrap local y un sistema de estilos propio en wwwroot/css/site.css. La navegación, tarjetas, formularios y mensajes comparten componentes visuales para que las rutas públicas, de usuario y de administración se perciban como un mismo producto.")
    add_table(doc, ["Criterio", "Implementación actual", "Criterio de mantenimiento"], [
        ("Intuitiva", "Barra con Inicio, Categorías y Destinos; buscador global; menús de usuario y administración; mensajes de éxito/error.", "Mantener nombres de acción claros, estado activo y acciones destructivas diferenciadas."),
        ("Responsiva", "Bootstrap con container/row/col, breakpoints md/lg, tarjetas que se apilan y navbar colapsable en móvil.", "Probar 320, 768 y 1280 px en cada pantalla nueva."),
        ("Accesible", "lang=es, enlace para saltar al contenido, labels asociados, aria-label/aria-controls, estados de validación y alt en imágenes.", "Verificar teclado, foco visible, contraste y que la información no dependa solo del color."),
        ("Cohesiva", "Variables de color, tipografías Poppins/Inter, botones y tarjetas reutilizables, parciales Razor.", "Reutilizar site.css y las parciales antes de crear estilos aislados.")
    ], [1.15, 3.05, 2.5])
    add_capture(doc, "04-catalogo-busqueda.jpg", "Figura 1. Captura actual del catálogo rediseñado con búsqueda, filtros y ordenamiento.", width=6.45)
    add_capture(doc, "22-inicio-movil.jpg", "Figura 2. La portada reorganiza el menú y el contenido para una pantalla móvil.", width=2.15)
    add_para(doc, "Puntos de atención para nuevas vistas: conservar encabezados jerárquicos, asociar cada input con un label, proporcionar mensajes de validación comprensibles, usar alt descriptivo en imágenes informativas y aria-hidden en iconos decorativos. El enlace externo de Google Fonts y Bootstrap Icons debe evaluarse si se despliega en un ambiente sin salida a Internet.")

    doc.add_heading("8 Seguridad y manejo de errores", level=1)
    add_table(doc, ["Riesgo o requisito", "Medida implementada"], [
        ("Autenticación", "Identity; mínimo 8 caracteres, dígito y minúscula; bloqueo de 10 minutos después de 5 intentos fallidos."),
        ("Autorización", "Roles Usuario y Administrador; [Authorize] en módulos privados y autorización por rol en Areas/Administracion."),
        ("Sesión", "Cookie de Identity con expiración deslizante de 90 minutos."),
        ("CSRF", "Tokens antifalsificación en formularios POST y validación en acciones de escritura."),
        ("XSS", "Razor codifica la salida; SanitizadorHtml elimina scripts, eventos y URLs no permitidas de iframes."),
        ("Imágenes", "JPG/JPEG/PNG, máximo 5 MB, tipo validado antes de leer bytes."),
        ("Headers", "nosniff, SAMEORIGIN, Referrer-Policy y Permissions-Policy; HTTPS redirection fuera de Development."),
        ("Errores", "UseExceptionHandler en producción, página de estado para 404/403 y no-store en páginas de error.")
    ], [1.55, 5.15])
    add_para(doc, "El contenido enriquecido se sanea al crear/actualizar y se vuelve a limpiar antes de mostrarlo. Solo se permiten iframes de Google Maps y YouTube con los atributos admitidos. Las excepciones de infraestructura deben observarse en logs, sin mostrar detalles internos al visitante.")

    doc.add_heading("9 Configuración del entorno", level=1)
    add_para(doc, "Requisitos mínimos: SDK .NET 10, una instancia SQL Server/LocalDB, editor con soporte C# y Razor, y acceso a NuGet durante el primer restore. La solución contiene TripSV/TripSV.csproj y TripSV.Pruebas/TripSV.Pruebas.csproj dentro de TripSV.slnx.")
    add_numbered(doc, [
        "Abrir una terminal en la raíz del repositorio.",
        "Configurar ConexionTripSV en User Secrets o en appsettings.json solo para desarrollo local.",
        "Restaurar y compilar con dotnet build TripSV.slnx.",
        "Crear/actualizar el esquema con dotnet ef database update --project TripSV.",
        "Cargar el catálogo con el script T-SQL si la base está vacía.",
        "Ejecutar dotnet run --project TripSV --launch-profile http y abrir http://localhost:5255."
    ])
    doc.add_page_break()
    add_table(doc, ["Usuario de desarrollo", "Contraseña", "Rol"], [("administrador", "Administrador123$", "Administrador"), ("visitante", "Visitante123$", "Usuario")], [2.2, 2.3, 2.2])
    add_para(doc, "Las credenciales anteriores son únicamente para demostración local. Deben cambiarse o eliminarse antes de publicar.", bold_lead="Las credenciales anteriores son únicamente para demostración local.")

    doc.add_heading("10 Despliegue", level=1)
    add_para(doc, "Se recomienda separar la publicación de la aplicación, la preparación de SQL Server y la carga de datos. Aunque el arranque aplica migraciones pendientes, en producción conviene ejecutar migraciones desde un pipeline con trazabilidad y respaldo previo.")
    add_numbered(doc, [
        "Crear la base de datos y un usuario de SQL Server con permisos mínimos.",
        "Definir la cadena de producción mediante secret manager o ConnectionStrings__ConexionTripSV.",
        "Publicar: dotnet publish TripSV/TripSV.csproj -c Release -o ./publish.",
        "Ejecutar dotnet ef database update --project TripSV desde el proceso controlado de despliegue.",
        "Cargar el catálogo inicial si el ambiente lo requiere, sin repetirlo sobre una base ya poblada.",
        "Hospedar detrás de IIS o reverse proxy con HTTPS, certificado válido, límites de carga y logs configurados.",
        "Persistir las claves de Data Protection en un directorio protegido o almacenamiento compartido si hay más de una instancia.",
        "Verificar las rutas públicas, login, autorización administrativa, imágenes, mensajes de error y respaldos."
    ])
    add_code(doc, "dotnet publish TripSV/TripSV.csproj -c Release -o ./publish\ndotnet TripSV.dll")
    add_para(doc, "El repositorio incluye Ejecutable/TripSV.zip como publicación Windows x64. Requiere ASP.NET Core Runtime 10, SQL Server y un appsettings/variable de entorno configurado para el ambiente.")

    doc.add_heading("11 Pruebas automatizadas", level=1)
    add_para(doc, "TripSV.Pruebas contiene 80 casos xUnit para calificaciones, comentarios y moderación, favoritos, itinerarios, categorías y destinos, validación de imágenes, saneamiento de HTML y reglas de seguridad de controladores. Los servicios usan una base SQLite en memoria para conservar relaciones, índices únicos y claves foráneas durante las pruebas.")
    add_table(doc, ["Suite", "Cobertura"], [
        ("CategoriasYSitiosServicioPruebas", "Búsqueda, categorías, destinos, duplicados y eliminaciones."),
        ("ComentariosServicioPruebas", "Publicación, respuestas, ocultamiento, eliminación y límites."),
        ("PuntuacionesServicioPruebas", "Rango 1–5, actualización y recálculo del promedio."),
        ("FavoritosServicioPruebas", "Agregar, quitar, duplicados y listado por usuario."),
        ("ItinerariosServicioPruebas", "Días, orden, movimientos, máximo por día y pertenencia."),
        ("Seguridad/ValidacionYSanitizacion", "Antiforgery, roles, imágenes, HTML permitido y URLs de iframe.")
    ], [2.8, 3.9])
    add_code(doc, "dotnet test")
    add_para(doc, "Antes de integrar cambios, ejecutar la suite completa, compilar en Release y comprobar también una base SQL Server con las migraciones. Las pruebas de SQLite no sustituyen la verificación del proveedor SQL Server en el ambiente destino.")

    doc.add_heading("12 Mantenimiento y extensión", level=1)
    add_table(doc, ["Tema", "Recomendación"], [
        ("Secretos", "Mover credenciales, cadenas y claves a secretos del ambiente; no versionar valores reales."),
        ("Migraciones", "Crear una migración por cambio de esquema y revisarla antes de aplicarla en producción."),
        ("Backups", "Respaldar SQL Server y validar restauración antes de migraciones destructivas."),
        ("Data Protection", "Persistir claves en servidores con más de una instancia para no invalidar cookies."),
        ("Contenido", "Conservar el saneamiento HTML y ampliar la lista de orígenes solo después de revisar seguridad."),
        ("Imágenes", "Mantener el límite de 5 MB, formatos permitidos y endpoint cacheado separado."),
        ("Pruebas", "Agregar casos de servicio y de autorización antes de incorporar un módulo nuevo.")
    ], [1.55, 5.15])
    add_para(doc, "Para agregar un módulo: crear entidad y configuración Fluent API, agregar DbSet y migración, implementar interfaz/servicio, registrarlo en Program.cs, crear controlador y ViewModels, construir vistas accesibles y añadir pruebas en TripSV.Pruebas. Mantener el Area para administración y [ValidateAntiForgeryToken] en toda acción POST.")
    add_capture(doc, "21-admin-moderacion.jpg", "Figura 3. Captura actual del panel de moderación rediseñado.", width=6.45)
    add_para(doc, "Checklist de entrega: compilar, ejecutar pruebas, aplicar migraciones en una base de prueba, validar secretos, comprobar HTTPS, revisar autorización por rol, verificar imágenes y contenido enriquecido, y confirmar que el catálogo esté cargado.")
    add_footer(doc, "Manual del Programador")
    doc.save(OUT / "Manual del Programador TripsSV.docx")


def build_user():
    doc = Document(); configure_doc(doc)
    add_cover(doc, "Manual del Usuario", "Guía de uso de Trips SV", "Visitantes, usuarios registrados y administradores.", "Explica la navegación pública, la búsqueda, la cuenta, favoritos, calificaciones, reseñas, itinerarios y herramientas administrativas con capturas reales de la aplicación.")
    add_toc(doc, [("1", "Antes de comenzar"), ("2", "Navegación general"), ("3", "Explorar destinos"), ("4", "Crear cuenta e iniciar sesión"), ("5", "Funciones del usuario registrado"), ("6", "Planificador de viajes"), ("7", "Funciones del administrador"), ("8", "Mensajes, errores y preguntas frecuentes")])

    doc.add_heading("1 Antes de comenzar", level=1)
    add_para(doc, "Trips SV permite descubrir destinos turísticos de El Salvador, consultar información, leer reseñas y organizar viajes. El catálogo es público; se requiere una cuenta para calificar, publicar, responder, guardar favoritos o crear itinerarios.")
    add_table(doc, ["Rol", "Funciones"], [
        ("Visitante", "Ver inicio, categorías, catálogo, filtros, detalle, calificaciones y reseñas."),
        ("Usuario registrado", "Todo lo anterior, más calificar, reseñar, responder, favoritos e itinerarios propios."),
        ("Administrador", "Todo lo anterior, más administrar destinos/categorías y moderar comentarios.")
    ], [1.65, 5.05])
    add_para(doc, "Se recomienda usar Edge, Chrome, Firefox o Safari actualizado. El sitio se adapta a computadora, tableta y teléfono. Para la demostración local se pueden usar administrador / Administrador123$ y visitante / Visitante123$; no son contraseñas para producción.")

    doc.add_heading("2 Navegación general", level=1)
    add_capture(doc, "01-inicio.jpg", "Figura 1. Página de inicio con navegación, buscador y acceso a la cuenta.")
    add_bullets(doc, [
        "Trips SV vuelve a la portada desde el logotipo.",
        "Inicio, Categorías y Destinos son las secciones públicas principales; la sección actual aparece resaltada.",
        "Buscar destinos permite lanzar una búsqueda desde cualquier página.",
        "Iniciar sesión y Registrarse abren las funciones de cuenta; con sesión iniciada aparece el menú del usuario.",
        "En un teléfono, el menú se agrupa en el botón de tres líneas de la esquina superior derecha."
    ])
    add_capture(doc, "22-inicio-movil.jpg", "Figura 2. Vista de la portada en un teléfono.", width=2.15)

    doc.add_heading("3 Explorar destinos", level=1)
    doc.add_heading("3.1 Explorar por categoría", level=2)
    add_numbered(doc, [
        "Desde la portada desplácese hasta la sección de tipos de viaje o pulse Categorías.",
        "Revise el nombre y el conteo de destinos de cada tarjeta.",
        "Pulse una tarjeta para ver sus destinos, ordenados de mejor a peor calificados.",
        "Abra un destino para ver su detalle."
    ])
    add_capture(doc, "02-inicio-categorias.jpg", "Figura 3. Categorías mostradas en la página de inicio.")
    add_capture(doc, "03-categorias.jpg", "Figura 4. Página completa de categorías.")
    doc.add_heading("3.2 Buscar y filtrar", level=2)
    add_numbered(doc, [
        "Escriba en el buscador un nombre, lugar, categoría o palabra de la descripción; por ejemplo lago, playa o Santa Ana.",
        "En Destinos combine Categoría, Departamento y Calificación mínima.",
        "Use Ordenar por para elegir mejor calificados, más votados, nombre o más recientes.",
        "Pulse Buscar. Use Limpiar filtros para regresar al catálogo completo."
    ])
    add_para(doc, "La búsqueda tolera tildes y mayúsculas. Si escribe varias palabras, se muestran destinos que coinciden con todas.")
    add_capture(doc, "04-catalogo-busqueda.jpg", "Figura 5. Catálogo rediseñado con búsqueda, filtros y ordenamiento.", width=6.45)
    add_capture(doc, "04-catalogo-busqueda.jpg", "Figura 6. Catálogo filtrado por texto y ordenamiento.")
    doc.add_heading("3.3 Ver el detalle", level=2)
    add_para(doc, "Abra cualquier tarjeta. El detalle muestra ubicación, categoría, promedio, favoritos, descripción, fotografía e información completa, que puede incluir horarios, recomendaciones, mapa o enlace externo de hoteles.")
    add_capture(doc, "05-detalle-sitio.jpg", "Figura 7. Detalle visto por un visitante.")
    doc.add_heading("3.4 Leer reseñas", level=2)
    add_para(doc, "Al final del detalle aparecen las reseñas de la comunidad y sus respuestas agrupadas. Para publicar o responder debe iniciar sesión.")
    add_capture(doc, "06-detalle-resenas.jpg", "Figura 8. Reseñas de la comunidad.")

    doc.add_heading("4 Crear una cuenta e iniciar sesión", level=1)
    doc.add_heading("4.1 Registrarse", level=2)
    add_numbered(doc, [
        "Pulse Registrarse en la barra superior.",
        "Escriba un usuario, correo, contraseña y confirmación.",
        "Use una contraseña de al menos 8 caracteres, con un dígito y una letra minúscula.",
        "Pulse Crear y, si todo es correcto, continúe con Iniciar sesión."
    ])
    add_capture(doc, "07-registro.jpg", "Figura 9. Formulario de registro.")
    doc.add_heading("4.2 Iniciar y cerrar sesión", level=2)
    add_numbered(doc, ["Pulse Iniciar sesión.", "Escriba usuario y contraseña.", "Marque Recordarme solo en un equipo personal.", "Pulse Entrar. Para salir, abra el menú con su nombre y pulse Cerrar sesión."])
    add_para(doc, "Después de cinco intentos fallidos la cuenta se bloquea durante diez minutos.")
    add_capture(doc, "08-iniciar-sesion.jpg", "Figura 10. Pantalla de inicio de sesión.")
    doc.add_heading("4.3 Recuperar contraseña", level=2)
    add_numbered(doc, ["En el inicio de sesión pulse ¿Olvidó su contraseña?", "Escriba su correo o usuario.", "Pulse Generar enlace de recuperación.", "Abra el enlace, defina la nueva contraseña y confirme el cambio."])
    add_capture(doc, "09-recuperar-contrasena.jpg", "Figura 11. Formulario de recuperación de contraseña.")
    doc.add_heading("4.4 Menú de usuario", level=2)
    add_para(doc, "Con sesión iniciada, su nombre aparece arriba a la derecha. El menú contiene Mis favoritos, Mis itinerarios y Cerrar sesión.")
    add_capture(doc, "12-menu-usuario.jpg", "Figura 12. Menú del usuario registrado.")

    doc.add_heading("5 Funciones del usuario registrado", level=1)
    add_capture(doc, "10-detalle-usuario.jpg", "Figura 13. Detalle con las tarjetas de favoritos, planificación y calificación.")
    doc.add_heading("5.1 Guardar en favoritos", level=2)
    add_numbered(doc, ["En el detalle pulse Guardar en mis favoritos.", "El botón cambia a Quitar de mis favoritos; úselo para retirarlo.", "Abra Mis favoritos desde el menú de usuario para consultar sus destinos guardados."])
    add_capture(doc, "13-mis-favoritos.jpg", "Figura 14. Listado de favoritos.")
    doc.add_heading("5.2 Calificar", level=2)
    add_numbered(doc, ["En la tarjeta Calificación elija de 1 a 5 estrellas.", "Pulse Calificar.", "Si ya calificó, el botón dice Actualizar mi calificación y reemplaza la anterior."])
    doc.add_heading("5.3 Publicar, responder o eliminar una reseña", level=2)
    add_numbered(doc, ["En Reseñas de la comunidad escriba su opinión y pulse Publicar reseña.", "Pulse Responder debajo de otra reseña, escriba y pulse Enviar respuesta.", "Puede eliminar sus propias reseñas y respuestas; el administrador también puede eliminarlas."])
    add_capture(doc, "11-detalle-publicar-resena.jpg", "Figura 15. Publicación y respuesta a una reseña.")

    doc.add_heading("6 Planificador de viajes", level=1)
    add_para(doc, "El planificador organiza los destinos que visitará cada día, con notas, orden y opción de impresión.")
    doc.add_heading("6.1 Crear un itinerario", level=2)
    add_numbered(doc, [
        "Abra Mis itinerarios y pulse Nuevo itinerario, o empiece desde el detalle con Crear itinerario con este destino.",
        "Escriba el nombre, seleccione fecha de inicio y defina de 1 a 15 días.",
        "Agregue notas opcionales y pulse Crear itinerario."
    ])
    add_capture(doc, "15-nuevo-itinerario.jpg", "Figura 16. Formulario de nuevo itinerario.")
    add_capture(doc, "14-mis-itinerarios.jpg", "Figura 17. Listado de itinerarios y su estado.")
    doc.add_heading("6.2 Organizar los días", level=2)
    add_numbered(doc, [
        "Pulse Abrir planificador en el itinerario.",
        "Elija un destino en Agregar un destino a este día, escriba una nota opcional y pulse Agregar.",
        "Use las flechas para cambiar el orden dentro del día.",
        "Use el selector Día para mover una visita.",
        "Pulse Quitar para retirarla y Cambiar día cuando la edición lo requiera.",
        "Pulse Imprimir para obtener una versión limpia."
    ])
    add_capture(doc, "16-planificador.jpg", "Figura 18. Planificador con destinos organizados por día.")
    add_para(doc, "Cada día admite hasta 10 destinos y no se puede repetir el mismo destino dentro del día. Si reduce la cantidad de días, primero quite o mueva las visitas de los días que desaparecerán.")

    doc.add_heading("7 Funciones del administrador", level=1)
    add_para(doc, "Estas opciones solo aparecen para el rol Administrador. Un usuario normal que intente abrirlas directamente verá Acceso denegado.")
    add_capture(doc, "17-menu-administrar.jpg", "Figura 19. Menú Administrar para el rol Administrador.")
    doc.add_heading("7.1 Gestionar destinos", level=2)
    add_numbered(doc, [
        "Abra Administrar > Modificar sitios para ver la tabla.",
        "Pulse Editar para cambiar un destino o Ver para abrir su página pública.",
        "Para eliminar, marque uno o varios destinos y pulse Eliminar seleccionados; la operación también elimina contenido relacionado.",
        "Para agregar, abra Administrar > Agregar sitios, complete nombre, categoría, ubicación, descripción e imagen y pulse Agregar sitio.",
        "La imagen debe ser JPG, JPEG o PNG y no superar 5 MB. El editor permite contenido enriquecido; el sistema elimina código ejecutable y limita iframes a Maps/YouTube seguros."
    ])
    add_capture(doc, "18-admin-sitios.jpg", "Figura 20. Tabla de administración de destinos.")
    add_capture(doc, "19-admin-agregar-sitio.jpg", "Figura 21. Formulario para agregar un destino.")
    doc.add_heading("7.2 Gestionar categorías", level=2)
    add_numbered(doc, ["Abra Administrar > Categorías.", "Use Buscar categorías si necesita localizar una fila.", "Pulse Nueva categoría y complete nombre, descripción e imagen opcional.", "Use Editar para actualizar una categoría.", "Una categoría con destinos asociados no se puede eliminar."])
    add_capture(doc, "20-admin-categorias.jpg", "Figura 22. Administración de categorías.")
    doc.add_heading("7.3 Moderar comentarios", level=2)
    add_numbered(doc, ["Abra Administrar > Moderación.", "Busque por texto, usuario o sitio, o active Mostrar solo ocultos.", "Pulse Ocultar para retirar un comentario sin borrarlo; sus respuestas también se ocultan.", "Pulse Mostrar para restaurarlo.", "Pulse Eliminar para borrarlo definitivamente junto con sus respuestas."])
    add_capture(doc, "21-admin-moderacion.jpg", "Figura 23. Panel de moderación.")

    doc.add_heading("8 Mensajes, errores y preguntas frecuentes", level=1)
    add_table(doc, ["Situación", "Qué hacer"], [
        ("Debe iniciar sesión", "Inicie sesión; el sistema intenta devolverlo a la página original."),
        ("No hay resultados", "Use Limpiar filtros o reduzca el texto de búsqueda."),
        ("Acceso denegado", "Solicite el rol Administrador al responsable del sistema."),
        ("Imagen no válida", "Use JPG/JPEG/PNG de hasta 5 MB."),
        ("No aparecen categorías o sitios", "Solicite cargar el script de datos inicial."),
        ("La aplicación no inicia", "Verifique SQL Server/LocalDB y la cadena ConexionTripSV.")
    ], [2.0, 4.7])
    add_capture(doc, "23-error-404.jpg", "Figura 24. Página mostrada cuando la dirección no existe.")
    doc.add_heading("Preguntas frecuentes", level=2)
    add_bullets(doc, [
        "¿Necesito cuenta para ver destinos? No. La cuenta solo se requiere para participar y usar favoritos o itinerarios.",
        "¿Puedo cambiar mi calificación? Sí, vuelva a elegir las estrellas y pulse Actualizar mi calificación.",
        "¿Otros usuarios pueden ver mis favoritos o itinerarios? No; son privados. Solo se muestra el total de favoritos de un destino.",
        "¿Por qué desapareció una reseña? Puede haber sido ocultada por un administrador.",
        "¿Por qué mi cuenta está bloqueada? Cinco intentos fallidos activan un bloqueo de diez minutos."
    ])
    add_para(doc, "Al terminar una sesión administrativa, cierre sesión. No comparta credenciales de demostración ni introduzca datos personales en un ambiente de prueba.")
    add_footer(doc, "Manual del Usuario")
    doc.save(OUT / "Manual del Usuario TripsSV.docx")


if __name__ == "__main__":
    make_figures()
    build_programmer()
    build_user()
    print("Manuales generados en", OUT)
